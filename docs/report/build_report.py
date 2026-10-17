"""Builds the shared draft design report (docs/report/Thought_Dock_Design_Report_DRAFT.docx).
Run: uv run --with python-docx python docs/report/build_report.py
Text marked [[like this]] is a placeholder the team must fill in; it renders highlighted yellow.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).parent
FIG = HERE / "figures"
SRC = HERE / "source"
OUT = HERE / "Thought_Dock_Design_Report_DRAFT.docx"
LIVE = "https://thought-dock-peach.vercel.app"

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(1)
sec.top_margin = sec.bottom_margin = Inches(1)

st = doc.styles["Normal"]
st.font.name = "Arial"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    h = doc.styles[name]
    h.font.name = "Arial"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = RGBColor(0x1D, 0x1D, 0x1F)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

TOKEN = re.compile(r"(\[\[.*?\]\]|\*\*.*?\*\*|\*[^*]+?\*)")
fig_n = 0
tab_n = 0


def runs(p, text, size=None, color=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("[["):
            r = p.add_run(part[2:-2])
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part.startswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*") and part.endswith("*"):
            r = p.add_run(part[1:-1])
            r.italic = True
        else:
            r = p.add_run(part)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def para(text, style=None, align=None, size=None, color=None, after=None):
    p = doc.add_paragraph(style=style)
    runs(p, text, size, color)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        para(it, style)


def h(text, level):
    return doc.add_heading(text, level)


def caption(text):
    p = para(text, size=9.5, color=RGBColor(0x42, 0x42, 0x45), after=12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def figure(path, cap, width=6.5):
    global fig_n
    fig_n += 1
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(f"**Figure {fig_n}.** {cap}")
    return fig_n


def figure_pair(a, b, cap, widths=(2.3, 4.1)):
    global fig_n
    fig_n += 1
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, path, w in zip(t.rows[0].cells, (a, b), widths):
        cell.paragraphs[0].add_run().add_picture(str(path), width=Inches(w))
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(f"**Figure {fig_n}.** {cap}")


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(header, rows, cap, widths=None, size=9.5):
    global tab_n
    tab_n += 1
    p = para(f"**Table {tab_n}.** {cap}", size=9.5, after=4)
    p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, head in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        runs(c.paragraphs[0], f"**{head}**", size)
        shade(c, "E8EEF7")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            runs(cells[i].paragraphs[0], str(val), size)
    if widths:
        t.autofit = False
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[i] * 1440)))
        for i, w in enumerate(widths):
            t.columns[i].width = Inches(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ============================== TITLE PAGE ==============================
for _ in range(5):
    doc.add_paragraph()
para("**Thought Dock / DeskOS**", align=WD_ALIGN_PARAGRAPH.CENTER, size=28)
para("Capture Now, Organize Later: An Accessible Thought Workspace for Overwhelmed Students",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=14, color=RGBColor(0x42, 0x42, 0x45))
doc.add_paragraph()
para("**Final Design Report**", align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
para("Introduction to Engineering · Mini-Hackathon Project", align=WD_ALIGN_PARAGRAPH.CENTER)
para("Theme: “Making an Accessible World: Every Dog Has Its Day”", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
para("Prepared by: [[Student full name]]", align=WD_ALIGN_PARAGRAPH.CENTER)
para("Team: Luke [[last name]], Nikhilesh Suravarjjala, Maryam [[last name]]", align=WD_ALIGN_PARAGRAPH.CENTER)
para("Instructor: Jared Gochuico, P.E.", align=WD_ALIGN_PARAGRAPH.CENTER)
para("[[Submission date]]", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
note = para("[[TEAM NOTE, DELETE BEFORE SUBMITTING: This is a shared draft. The course requires every "
            "student to submit an individual report, so each person must make a copy and revise it into "
            "their own version. Everything highlighted yellow needs a real human answer. Do not fill in "
            "test results or matrix scores that did not actually happen.]]", size=9.5)
page_break()

# ============================== 1. EXECUTIVE SUMMARY ==============================
h("1. Executive Summary/Abstract", 1)
para("Students collect assignments, deadlines, links, ideas, and worries throughout the day, and most productivity "
     "tools ask them to categorize each item at the exact moment their attention is most limited. Working memory "
     "holds only about three to five items at once (Cowan, 2001), and unfinished goals continue to intrude on "
     "attention until a plan for them exists (Masicampo & Baumeister, 2011). The goal of this project was to design "
     "a software-only workspace that separates *capturing* a thought from *organizing* it, for students who "
     "experience cognitive overload, including students with attention or executive-function difficulties.")
para("The design criteria required capture in one tap with zero required decisions, fully reversible and "
     "keyboard-operable processing, text contrast of at least 4.5:1 (WCAG 2.2 AA), no cost and no account, and a "
     "reliable live demonstration. Three concepts were developed: a physical triangular NFC dock, a software-only "
     "Capture Mode → DeskOS application, and an AI-first automatic organizer. The software-only concept was selected "
     "because it kept every decision with the user at no cost while remaining feasible within one week "
     "([[TEAM: add the final weighted decision-matrix totals from Table 7]]).")
para(f"The final design is a working web application ({LIVE}). On a phone, the student types, speaks (with a live "
     "transcript), or photographs a thought and taps Save. At a desk, the student pairs by QR code without creating an "
     "account and processes captures one at a time with seven labeled, keyboard-accessible actions, each of which "
     "can be undone. A text classifier trained from scratch on 690 labeled examples suggests what each thought is "
     "(85.8% accuracy on held-out data) but never moves anything on its own. An automated contrast check passes in "
     "both light and dark themes after it identified and led to the correction of one failing color. "
     "[[TEAM: add one sentence on the usability test result.]] The main lessons were that automated checks catch "
     "accessibility problems that visual review misses, and that a model trained on synthetic data needs real user "
     "data before its accuracy can be trusted.")

# ============================== 2. DESIGN CRITERIA ==============================
h("2. Design Criteria and Goals", 1)
h("2.1 Problem", 2)
para("Existing note and task tools require a destination (folder, tag, project, or date) at the moment of capture. "
     "That decision competes for the same limited mental capacity the student is already using (Sweller, 1988). "
     "When capture is costly, thoughts are either lost or held in mind, and switching tasks to organize them leaves "
     "attention behind on the previous task (Leroy, 2009). Interrupted work is completed faster but with "
     "measurably higher stress (Mark et al., 2008).")
h("2.2 End user and context", 2)
para("The end user is an overwhelmed student: commuting community-college students who work part-time, transfer "
     "applicants juggling deadlines, and students with ADHD-like attention patterns. An estimated 2–8% of college "
     "students report clinically significant ADHD symptoms, and those symptoms are consistently linked to academic "
     "difficulty (DuPaul et al., 2009). Executive functions such as working memory and inhibition are also "
     "impaired by stress and poor sleep, which affects nearly all students at some point (Diamond, 2013). The "
     "design therefore targets cognitive load in general rather than any single diagnosis. Externalizing thoughts "
     "to a trusted tool, known as cognitive offloading, is an established strategy for reducing mental demand "
     "(Risko & Gilbert, 2016).")
para("Two usage contexts drive the design: *on the go* (a phone, one hand, a few seconds, in class or on a bus) and "
     "*at a desk* (a laptop, a deliberate focus session, a keyboard).")
h("2.3 Constraints", 2)
bullets([
    "One-week build window; team of three; $0 budget.",
    "Software only: no required hardware. (The NFC dock was explored as a concept but is not required.)",
    "Must work without paid AI APIs and without user accounts.",
    "Must be demonstrated live within a 15-minute presentation.",
    "Must meet WCAG 2.2 Level AA where applicable (World Wide Web Consortium [W3C], 2023).",
])
h("2.4 Design criteria", 2)
table(["ID", "Criterion", "Measurable target", "How verified", "Result"], [
    ["C1", "Capture effort", "Save in 1 tap after typing; 0 required fields besides the thought", "UI inspection (Figure 7)", "Met"],
    ["C2", "Capture-to-desk time", "Saved item visible on desk within one 4 s polling interval", "Design calculation (Section 4.6)", "Met by design; [[measure on a phone]]"],
    ["C3", "Reversibility", "Every processing action undoable; no hard delete outside Archive", "Automated action checks", "Met"],
    ["C4", "Keyboard access", "Every desk action has a visible labeled button and a single-key shortcut", "Inspection + keyboard walkthrough", "Met; [[human walkthrough]]"],
    ["C5", "Text contrast", "≥ 4.5:1 for all text tokens in light and dark themes", "npm run check:contrast", "Met (lowest 4.66:1)"],
    ["C6", "Motion", "All movement removed under reduced-motion setting", "Code review + manual check", "Met in code; [[manual check]]"],
    ["C7", "Target size", "≥ 44 px on phone; ≥ 24 px on desktop (WCAG 2.5.8)", "Measured in browser", "48 px phone; 40 px desktop"],
    ["C8", "Cost and access", "$0 to use; no account required", "Hosting on free tiers; QR pairing", "Met"],
    ["C9", "User control", "Suggestions never change data; shown only at ≥ 0.45 confidence", "Code review", "Met"],
    ["C10", "Gentle language", "No streaks, scores, or red overdue alerts", "UI inspection", "Met"],
], "Design criteria with measurable targets.", widths=(0.4, 1.1, 2.1, 1.5, 1.4))
h("2.5 Goals", 2)
bullets([
    "**Primary:** reduce the effort of capturing a thought to a single action, and make later processing calm, "
    "one item at a time, and fully reversible.",
    "**Secondary:** keep all data under the user's control, run without cost, and qualify the from-scratch model for "
    "the course's AI award.",
])

# ============================== 3. BRAINSTORM ==============================
h("3. Brainstorm/Concepts", 1)
h("3.1 Engineering design process", 2)
table(["Step", "What the team did"], [
    ["Define", "Wrote the problem statement and a point-of-view: a busy student needs to get a thought out of their head the moment it appears, without deciding what it is."],
    ["Research", "Reviewed literature on working memory, cognitive offloading, decision fatigue, attention residue, and ADHD in college students; studied WCAG 2.2."],
    ["Ideate", "Generated 20 ideas in five clusters (Table 3), then combined them into three distinct concepts."],
    ["Select", "Compared concepts in a weighted decision matrix (Section 5)."],
    ["Prototype", "Hand drawing and render of Concept 1; wireframes, then a coded prototype of Concept 2."],
    ["Test", "Automated checks (contrast, type, lint, build, model tests, production smoke test); usability test protocol [[results]]."],
    ["Iterate", "Fixed a failing contrast color and a keyboard-shortcut conflict; [[changes from usability test]]."],
], "How the engineering design process was applied.", widths=(1.0, 5.5))
h("3.2 “How might we” questions", 2)
bullets([
    "How might a student save a thought in under five seconds, one-handed, without choosing anything?",
    "How might sorting feel like a short task with an end, instead of a pile that keeps growing?",
    "How might the tool offer structure (“this looks like a task”) without taking the decision away?",
])
h("3.3 Brainstorm", 2)
table(["Capture", "Sorting", "Hardware", "Help and AI", "Calm"], [
    ["One text field, nothing else", "Sort one thought at a time", "NFC dock on the desk", "Suggest a type; user confirms", "Undo on every action"],
    ["Voice note button", "Turn a thought into a task", "Tap phone to enter Desk Mode", "Auto-file everything with AI", "No red overdue badges"],
    ["Photo of a whiteboard", "Turn it into a calendar event", "Physical inbox light", "Daily digest email", "Pair with a QR code, no login"],
    ["Works offline, syncs later", "Spaces per class", "", "", "Streaks and points (rejected)"],
    ["Lock-screen widget", "Link related thoughts", "", "", ""],
], "Brainstorm wall, grouped by cluster.")
h("3.4 Three concepts", 2)
figure(FIG / "fig2-concepts.png", "Concept sketches for the three alternatives. Concept 2 was selected.")
para("**Concept 1: Triangular Prism NFC Dock.** A 170 × 85 × 68 mm triangular prism sits on the desk. Its 39° "
     "face holds a MagSafe puck and an NFC tag. Setting the phone on the dock would charge it and switch the app into "
     "Desk Mode, turning the move from capture to processing into a physical ritual (Figures 2 and 3). "
     "*Strengths:* tactile, memorable, and it makes the transition deliberate. *Weaknesses:* it costs money to "
     "build, it works only at one desk, and requiring an object to reach Desk Mode is itself an access barrier.")
para("**Concept 2: Capture Mode → DeskOS (selected).** Software only. The phone screen does one thing, capture, and "
     "a desktop workspace is used for sorting when the student chooses to sit down. *Strengths:* free, runs on "
     "devices students already own, and keeps every decision with the user. *Weaknesses:* less physical delight, "
     "and the transition to Desk Mode relies on habit rather than an object.")
para("**Concept 3: AI-first automatic organizer.** All thoughts go into one box and a language model files them "
     "automatically into tasks, events, and folders. *Strengths:* the least effort for the user when it works. "
     "*Weaknesses:* it removes user control, can silently misfile items, requires paid APIs, and sends personal "
     "thoughts to a third party.")
figure(SRC / "ai-ref-image.png", "Concept 1 reference render: triangular prism dock with NFC chip and MagSafe puck, shown with a phone. "
       "AI-generated visualization for concept communication only, not a photograph of a built part.", 5.2)

# ============================== 4. TECHNICAL ==============================
h("4. Technical and Visual Diagrams, Documentation", 1)
h("4.1 Concept 1 dimensioned drawing", 2)
para("Figure 3 is the team's hand-drawn, fully dimensioned multiview drawing of the Concept 1 dock (front, top, "
     "right side, and Section A-A, plus an isometric reference view). All dimensions are in millimetres. Overall "
     "size is 170 (W) × 85 (D) × 68 (H) with a 39° sloped face. Features: a Ø56 × 6 deep MagSafe puck recess, a "
     "30 × 30 × 2 deep NFC chip pocket, a 6 mm wide internal cable channel, and a Ø8 rear cable exit hole.")
figure(SRC / "prism-dock-p1.png", "Concept 1, Thought Dock / DeskOS Triangular Prism Dock: dimensioned multiview drawing (mm). "
       "[[TEAM: check the drawing against the notes before submitting, see Section 7.]]", 5.6)
para("**Check of the 39° face:** for a right-triangle profile, tan θ = height ÷ run, so the run for a 68 mm height "
     "is 68 ÷ tan 39° = 68 ÷ 0.810 = 84.0 mm. This is consistent with the 85 mm depth, confirming that the "
     "angle and the overall dimensions agree.")
para("**Material estimate:** the solid volume of the prism is ½ × 85 × 68 × 170 = 491,300 mm³ ≈ 491 cm³. At "
     "1.24 g/cm³ for PLA, a solid part would weigh about 609 g. Printed with walls and roughly 20% infill, an "
     "effective fill of about 25% gives approximately 150 g of filament, or about [[$ per kg × 0.15]] in material, plus "
     "[[MagSafe puck price]] and [[NFC tag price]].")
h("4.2 Final design overview", 2)
figure_pair(FIG / "fig1a-capture.png", FIG / "fig1b-desk.png",
            "Final design (Concept 2). Left: Capture screen on a phone. Right: DeskOS on a laptop, showing the Inbox, the selected "
            "thought, and the labeled processing actions with their keyboard keys.")
h("4.3 User flow", 2)
figure(FIG / "fig7-user-flow.png", "User flow from capture on the phone to processing on the desk.")
h("4.4 System architecture and data model", 2)
figure(FIG / "fig6-architecture.png", "System architecture. The phone and desk exchange small JSON operations through Next.js API routes "
       "on Vercel; data is stored in Upstash Redis; the classifier runs entirely in the browser.")
table(["Field", "Type", "Purpose"], [
    ["id", "string", "Random identifier"],
    ["text", "string", "What was captured, verbatim (never edited in place)"],
    ["status", "inbox | kept | archived", "Inbox until processed; Archive is reversible"],
    ["kind", "task | event | idea | resource | note", "Set by the user when processing"],
    ["spaceId", "string", "Personal space such as Classes or Engineering"],
    ["due / at", "date / date-time", "Task due date or event start"],
    ["links", "string[]", "Connected thoughts (kept two-way)"],
    ["media", "audio | image", "Voice note or photo, stored separately"],
], "Core data model for one captured thought.", widths=(1.1, 2.2, 3.2))
para("Every change is sent as a small operation (put or delete one item). Each item is a separate field in a Redis "
     "hash, so a phone saving one thought and a desk editing another never overwrite each other. Changes appear "
     "immediately on screen, are queued in browser storage, and are retried if the connection drops.")
h("4.5 Dimensioned screen layouts of the final design", 2)
para("Because the final design is software, its controlling dimensions are screen dimensions. Figures 7 and 8 give "
     "them in CSS pixels (1 px = 1/96 in at 100% zoom), measured in the deployed application. Key values: the phone "
     "Save button is 94 × 48 px (above the 44 px minimum touch target); the desktop uses a 248 px sidebar, a 360 px "
     "list, and a flexible detail column; processing buttons are 40 px tall.")
figure(FIG / "fig3-capture-layout.png", "Capture screen, dimensioned layout (390 × 844 px phone viewport).", 5.6)
figure(FIG / "fig4-desk-layout.png", "DeskOS, dimensioned layout (1440 × 900 px desktop viewport).")
figure(FIG / "fig8-sort.png", "Sort view: one thought at a time, a suggestion chip from the classifier, and seven labeled actions with keys.")
h("4.6 Engineering calculations", 2)
para("**Voice note storage.** Audio is recorded at 24 kbit/s for at most 60 s: 24 kbit/s × 60 s = 1,440 kbit = "
     "180 kB. Base64 encoding adds 4/3 overhead: 180 × 4/3 = 240 kB, below the 900 kB per-item limit by a factor "
     "of 3.75. The 60 s limit was chosen to keep notes short and leave margin.")
para("**Capture-to-desk delay.** The desk polls every T = 4 s, so the expected wait after a save is T/2 = 2 s and "
     "the worst case is 4 s plus one network round trip ([[TEAM: measure on a real phone]]).")
para("**Server load.** One open desk makes 3,600 ÷ 4 = 900 requests per hour, each using 3 database commands: "
     "2,700 commands per hour. Polling pauses whenever the tab is hidden.")
para("**Contrast.** WCAG contrast ratio = (L₁ + 0.05) ÷ (L₂ + 0.05), where L is relative luminance. The original "
     "accent #0071E3 on the #F5F5F7 background measured 4.31:1, below the 4.5:1 minimum. Changing it to #0066CC "
     "gives 5.11:1, which passes.")
para("**Model accuracy uncertainty.** Held-out accuracy p = 0.858 on n = 141 examples. Standard error = "
     "√(p(1 − p)/n) = √(0.858 × 0.142 / 141) = 0.0294. The 95% interval is 0.858 ± 1.96 × 0.0294 = 80.0% to 91.6%.")
h("4.7 Classifier trained from scratch", 2)
para("The suggestion model is a softmax (multinomial logistic) regression written and trained from scratch in "
     "TypeScript, with no pre-trained weights or ML libraries. Features are word unigrams and bigrams plus six "
     "engineered signals (URL, time, date word, imperative verb, question mark, length). It was trained by "
     "mini-batch gradient descent (60 epochs, learning rate 0.5, L2 regularization) on 690 labeled thoughts and "
     "evaluated on a stratified 20% held-out split. The weights (42 KB) ship with the app, so classification "
     "happens on the device and no thought leaves it to be classified. The training data was written by an AI "
     "assistant and labeled with a written guide; real student captures are needed before accuracy can be "
     "trusted for real users.")
table(["Model", "Held-out accuracy", "Macro-F1"], [
    ["Softmax regression (shipped)", "85.8%", "85.7%"],
    ["Naive Bayes (baseline)", "81.6%", "81.0%"],
    ["Keyword rules (baseline)", "77.3%", "78.0%"],
], "Classifier comparison on the same 141 held-out examples. Five-fold cross-validation of the shipped model: 84.9% ± 1.6%.",
    widths=(2.8, 1.8, 1.4))
h("4.8 Accessibility design", 2)
table(["Feature", "Implementation", "WCAG 2.2 criterion"], [
    ["Contrast", "All text tokens ≥ 4.5:1 in both themes, checked automatically", "1.4.3 Contrast (Minimum)"],
    ["Color is never the only signal", "Every thought type shows an icon and a word", "1.4.1 Use of Color"],
    ["Keyboard operation", "Every action has a visible button and a single key", "2.1.1 Keyboard"],
    ["Visible focus", "3 px accent outline on every control", "2.4.7 Focus Visible"],
    ["Reduced motion", "System setting and in-app toggle remove movement", "2.3.3 Animation from Interactions"],
    ["Target size", "48 px phone controls, 40 px desktop controls", "2.5.8 Target Size (Minimum)"],
    ["Status messages", "Toasts and sync status announced politely", "4.1.3 Status Messages"],
    ["No sign-in", "QR / 6-character code pairing, no password", "3.3.8 Accessible Authentication"],
], "Accessibility features mapped to WCAG 2.2 success criteria (W3C, 2023).", widths=(1.5, 3.1, 1.9))

# ============================== 5. ALTERNATIVES ==============================
h("5. Alternatives Analysis (Decision Matrix)", 1)
para("The three concepts were compared with a weighted decision matrix. Weights sum to 1.0, each concept is scored "
     "out of 10 per criterion, and the weighted score is weight × score. [[TEAM: agree on the weights first, then "
     "score each concept together; issue #17. These must be the team's own numbers.]]")
crit = ["Accessibility", "Mental effort at capture", "Cost", "Feasibility in one week", "Ease of demonstration", "User control", "Distinctiveness"]
table(["Criterion", "Weight", "C1 /10", "C1 wtd", "C2 /10", "C2 wtd", "C3 /10", "C3 wtd"],
      [[c, "[[  ]]", "[[  ]]", "", "[[  ]]", "", "[[  ]]", ""] for c in crit] + [["**Total**", "**1.0**", "", "[[  ]]", "", "[[  ]]", "", "[[  ]]"]],
      "Weighted decision matrix. C1 = Prism NFC Dock, C2 = Capture → DeskOS, C3 = AI-first organizer.",
      widths=(1.9, 0.65, 0.6, 0.6, 0.6, 0.6, 0.6, 0.6), size=9)
table(["Criterion", "C1 Prism NFC Dock", "C2 Capture → DeskOS", "C3 AI-first organizer"], [
    ["Accessibility", "Requires a physical object at a desk", "Runs on any phone and browser", "Accessible UI possible; errors are hard to notice"],
    ["Mental effort at capture", "Low once docked, but the dock must be nearby", "Lowest: one field, no choices", "Low at capture, high when checking its filing"],
    ["Cost", "Printing, MagSafe puck, NFC tags", "Free tiers only", "Ongoing API cost per thought"],
    ["Feasibility in one week", "Hardware plus software", "Software the team can finish and polish", "Depends on prompt quality and API limits"],
    ["Ease of demonstration", "Needs the dock in the room", "Two screens, live, on a laptop", "Results vary from run to run"],
    ["User control", "Full", "Full: suggestions only, undo everywhere", "Low: the model decides first"],
    ["Distinctiveness", "Most distinctive object", "Distinct split between capture and sorting", "Many similar products exist"],
], "Qualitative comparison used to inform the scores.", widths=(1.3, 1.7, 1.75, 1.75), size=9)
para("**Selection.** Concept 2 was selected. It is the only option that is free, requires no specific object or "
     "location, and keeps every decision with the user, while still being buildable and demonstrable within the "
     "time available. Concept 1's strongest idea, a deliberate switch into Desk Mode, was kept in software as a "
     "separate desk workspace and the Sort view. Concept 3's strongest idea, help with categorizing, was kept as an "
     "on-device suggestion that the user must confirm. [[TEAM: state the final weighted totals and the margin by "
     "which Concept 2 won.]]")

# ============================== 6. PROPOSED SOLUTION ==============================
h("6. Proposed Solution", 1)
para(f"The final design is the Thought Dock / DeskOS web application, deployed at {LIVE}. A student pairs a phone "
     "to a laptop once by scanning a QR code. During the day, the phone shows a single field with Voice, Photo, "
     "and Save. Later, the desk shows an Inbox; **Sort** presents one thought at a time with seven actions: "
     "Make a Task (T), Make an Event (E, with a calendar file for any calendar app), Move to Space (M), Develop (D), "
     "Connect (C), Keep (K), and Archive (A). Every action can be undone with ⌘Z or U.")
table(["Check", "Result"], [
    ["Type check, lint, production build", "Pass"],
    ["Classifier tests", "16 of 16 checks pass"],
    ["Contrast check (16 token pairs × 2 themes)", "All pass; lowest 4.66:1"],
    ["Production smoke test (create, pair, save, read, reject malformed data)", "Pass"],
    ["Automated desk-action checks (headless browser)", "17 of 17 pass"],
    ["Real-phone voice, transcript, and camera", "[[Not yet tested]]"],
    ["Screen reader (VoiceOver) walkthrough", "[[Not yet tested]]"],
    ["Usability test with one student", "[[Not yet run, see docs/USABILITY_TEST.md]]"],
], "Verification results.", widths=(4.2, 2.3))
h("6.1 Usability test", 2)
para("Task given to a student outside the team, with minimal help: “Capture an application deadline and useful "
     "link on your phone. Later, turn the deadline into a task and save the link in an Engineering space.” "
     "[[TEAM: record where the participant hesitated, what confused them, what was easiest, what added friction, "
     "what they expected next, and at least one change made because of it.]]")

# ============================== 7. LESSONS LEARNED ==============================
h("7. Lessons Learned", 1)
bullets([
    "**Automated checks catch what visual review misses.** The accent blue looked fine but measured 4.31:1 against the "
    "page background. A scripted contrast check found it, and the color was darkened to pass at 5.11:1. The check now "
    "runs with every test.",
    "**Specifications need their own review.** The keyboard plan assigned the K key to both “Keep” and “previous item”. "
    "The conflict was found during the build and resolved by keeping K for Keep and using the arrow keys to navigate.",
    "**Parallel work needs clear ownership.** The app was built in four parallel streams (interface, classifier, "
    "portfolio page, landing page) with separate files for each. The only merge conflict happened in the one shared "
    "file, the icon set.",
    "**Synthetic data is a starting point, not evidence.** The classifier's 85.8% accuracy was measured on examples "
    "written by an AI assistant. Real student captures are needed before claiming real-world accuracy.",
    "**Hardware features cannot be verified in simulation.** Microphone, speech recognition, and camera behavior "
    "could not be tested in a headless browser and still require real phones.",
    "**Drawings must agree with each other.** [[TEAM: check Figure 3: the right side view shows a right-triangle "
    "profile while Section A-A and the isometric show an isosceles prism, the front view is drawn as a trapezoid, "
    "and “TOP VIE” is cut off. Decide which profile is correct and redraw if needed.]]",
    "[[TEAM: one lesson from the usability test.]]",
    "[[EACH STUDENT: one personal reflection on what to do differently next time.]]",
])

# ============================== REFERENCES ==============================
h("References", 1)
refs = [
    "Cowan, N. (2001). The magical number 4 in short-term memory: A reconsideration of mental storage capacity. *Behavioral and Brain Sciences, 24*(1), 87–114. https://doi.org/10.1017/S0140525X01003922",
    "Diamond, A. (2013). Executive functions. *Annual Review of Psychology, 64*, 135–168. https://doi.org/10.1146/annurev-psych-113011-143750",
    "DuPaul, G. J., Weyandt, L. L., O'Dell, S. M., & Varejao, M. (2009). College students with ADHD: Current status and future directions. *Journal of Attention Disorders, 13*(3), 234–250. https://doi.org/10.1177/1087054709340650",
    "Leroy, S. (2009). Why is it so hard to do my work? The challenge of attention residue when switching between work tasks. *Organizational Behavior and Human Decision Processes, 109*(2), 168–181. https://doi.org/10.1016/j.obhdp.2009.04.002",
    "Mark, G., Gudith, D., & Klocke, U. (2008). The cost of interrupted work: More speed and stress. In *Proceedings of the SIGCHI Conference on Human Factors in Computing Systems* (pp. 107–110). ACM. https://doi.org/10.1145/1357054.1357072",
    "Masicampo, E. J., & Baumeister, R. F. (2011). Consider it done! Plan making can eliminate the cognitive effects of unfulfilled goals. *Journal of Personality and Social Psychology, 101*(4), 667–683. https://doi.org/10.1037/a0024192",
    "Risko, E. F., & Gilbert, S. J. (2016). Cognitive offloading. *Trends in Cognitive Sciences, 20*(9), 676–688. https://doi.org/10.1016/j.tics.2016.07.002",
    "Sweller, J. (1988). Cognitive load during problem solving: Effects on learning. *Cognitive Science, 12*(2), 257–285. https://doi.org/10.1207/s15516709cog1202_4",
    "Vohs, K. D., Baumeister, R. F., Schmeichel, B. J., Twenge, J. M., Nelson, N. M., & Tice, D. M. (2008). Making choices impairs subsequent self-control: A limited-resource account of decision making, self-regulation, and active initiative. *Journal of Personality and Social Psychology, 94*(5), 883–898. https://doi.org/10.1037/0022-3514.94.5.883",
    "World Wide Web Consortium. (2023). *Web Content Accessibility Guidelines (WCAG) 2.2* (W3C Recommendation, 5 October 2023). https://www.w3.org/TR/2023/REC-WCAG22-20231005/",
]
for r in refs:
    p = para(r)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)

# ============================== APPENDICES ==============================
h("Appendix A. Use of AI tools", 1)
para("Claude Code (Anthropic) was used as a design and software-building assistant: it proposed design directions, "
     "wrote the application code, generated the screen-layout and system diagrams, trained the classifier from the "
     "team's specification, and drafted this report. The Concept 1 render (Figure 2) is AI-generated. The team set "
     "the problem, chose the concept, made the product decisions recorded in the project's decision log, produced "
     "the hand drawing (Figure 3), [[scored the decision matrix, ran the usability test,]] and reviewed and revised "
     "this report. [[EACH STUDENT: adjust this paragraph to describe your own contribution accurately.]]")
h("Appendix B. Project links", 1)
bullets([
    f"Live application: {LIVE} (capture: /capture, desk: /desk)",
    f"Process portfolio: {LIVE}/jared · Launch video: {LIVE}/brag",
    "Source code: https://github.com/nikhilesh-s/thought-dock",
    "Project board: https://github.com/users/nikhilesh-s/projects/2",
])

doc.save(OUT)
print("saved", OUT, "figures:", fig_n, "tables:", tab_n)
